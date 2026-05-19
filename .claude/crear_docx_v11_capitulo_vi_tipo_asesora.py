from pathlib import Path
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Cm

BASE = Path(r"C:\Users\jquis\OneDrive\Escritorio\TESIS TOMAS\tesista")
SRC = BASE / "INVESTIGACION FINAL - PARA PRESENTAR_PREPARADO_FINAL_ASESORA_2026-05-18_v10.docx"
OUT = BASE / "INVESTIGACION FINAL - PARA PRESENTAR_CAPITULO_VI_TIPO_ASESORA_2026-05-18_v11.docx"

shutil.copy2(SRC, OUT)
doc = Document(OUT)

FONT = "Times New Roman"
SIZE = Pt(12)


def norm(text):
    return " ".join(text.split())


def element_text(el):
    texts = []
    for node in el.iter():
        if str(node.tag).endswith('}t') and node.text:
            texts.append(node.text)
    return " ".join(texts)


def apply_run(run, bold=None):
    run.font.name = FONT
    run.font.size = SIZE
    if bold is not None:
        run.bold = bold


def format_paragraph(p, heading=False, center=False, first_line=True, bold=None):
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif heading:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = None if heading or center or not first_line else Cm(1.25)
    for r in p.runs:
        apply_run(r, bold=bold if bold is not None else (True if heading else None))


def set_cell(cell, text, bold=False, center=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center or bold else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            apply_run(r, bold=bold)


def find_body_index_by_text(target):
    body = doc._body._element
    for idx, el in enumerate(list(body)):
        if target in norm(element_text(el)):
            return idx
    return None


def remove_body_range(start_idx, end_idx):
    body = doc._body._element
    for el in list(body)[start_idx:end_idx]:
        body.remove(el)


def insert_paragraph_at(index, text, *, heading=False, center=False, first_line=True, bold=None):
    p = doc.add_paragraph()
    p.text = text
    format_paragraph(p, heading=heading, center=center, first_line=first_line, bold=bold)
    body = doc._body._element
    body.remove(p._p)
    body.insert(index, p._p)
    return index + 1


def insert_table_at(index, rows, headers=True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            set_cell(table.cell(r, c), value, bold=(headers and r == 0), center=True)
    body = doc._body._element
    body.remove(table._tbl)
    body.insert(index, table._tbl)
    return index + 1


def add_caption(index, number, title):
    return insert_paragraph_at(index, f"Tabla {number}  | {title}", heading=True, first_line=False)


def add_note(index, text="Nota. Elaboración propia a partir de la estructura de contrastación estadística definida para la investigación."):
    return insert_paragraph_at(index, text, first_line=False)


def add_formula_block(index, formula, definitions):
    index = insert_paragraph_at(index, "Con esta información se procede a calcular el valor de t mediante la siguiente fórmula:")
    index = insert_paragraph_at(index, formula, center=True, first_line=False)
    index = insert_paragraph_at(index, "Donde:", first_line=False)
    for item in definitions:
        index = insert_paragraph_at(index, item, first_line=False)
    return index


def placeholder_execution_table(model_label, combined=False):
    if combined:
        return [
            ["Ejecución", "Modelo", "Seed", "Accuracy", "Interpretación"],
            ["1", "CNN/EfficientNetB0", "Por completar", "Por completar", "Por completar"],
            ["2", "CNN/EfficientNetB0", "Por completar", "Por completar", "Por completar"],
            ["3", "Swin-Tiny Transformer", "Por completar", "Por completar", "Por completar"],
            ["4", "Swin-Tiny Transformer", "Por completar", "Por completar", "Por completar"],
            ["...", "...", "...", "...", "..."],
        ]
    return [
        ["Ejecución", "Seed", "Configuración", "Accuracy", "Interpretación"],
        ["1", "Por completar", model_label, "Por completar", "Por completar"],
        ["2", "Por completar", model_label, "Por completar", "Por completar"],
        ["3", "Por completar", model_label, "Por completar", "Por completar"],
        ["4", "Por completar", model_label, "Por completar", "Por completar"],
        ["5", "Por completar", model_label, "Por completar", "Por completar"],
    ]


def one_sample_t_table():
    return [
        ["Indicador", "Valor"],
        ["Media (X̄)", "Por completar"],
        ["Desviación estándar (S)", "Por completar"],
        ["Observaciones (N)", "Por completar"],
        ["Valor de referencia (μ₀)", "0.75"],
        ["Grados de libertad", "N - 1"],
        ["Estadístico t", "Por completar"],
        ["p valor", "Por completar"],
        ["Nivel de significancia (α)", "0.05"],
        ["Decisión", "Por completar"],
    ]


def paired_t_table():
    return [
        ["Indicador", "Valor"],
        ["Media de diferencias (D̄)", "Por completar"],
        ["Desviación estándar de diferencias (S_D)", "Por completar"],
        ["Observaciones pareadas (N)", "Por completar"],
        ["Grados de libertad", "N - 1"],
        ["Estadístico t", "Por completar"],
        ["p valor", "Por completar"],
        ["Nivel de significancia (α)", "0.05"],
        ["Decisión", "Por completar"],
    ]


start = find_body_index_by_text("CAPÍTULO VI: RESULTADOS")
end = find_body_index_by_text("CAPÍTULO VII: DISCUSIÓN")
if start is None or end is None or end <= start:
    raise RuntimeError("No se pudo ubicar el bloque del Capítulo VI.")
remove_body_range(start, end)
idx = start

# Capítulo VI tipo asesora.
idx = insert_paragraph_at(idx, "CAPÍTULO VI: RESULTADOS", heading=True, center=True, first_line=False, bold=True)
idx = insert_paragraph_at(idx, "PRESENTACIÓN DE RESULTADOS Y ANÁLISIS", heading=True, center=True, first_line=False, bold=True)
idx = insert_paragraph_at(idx, "6.1. CONTRASTACIÓN DE HIPÓTESIS", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "6.1.1. Hipótesis de investigación", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "La investigación plantea que los algoritmos de Deep Learning presentan un desempeño eficiente en la identificación de variedades de papas nativas en la comunidad de Santa Carmen. Para sustentar esta afirmación se considera como evidencia principal el Accuracy obtenido en ejecuciones completas del entrenamiento, complementado con precision, recall, F1-score, balanced accuracy y matrices de confusión.")
idx = insert_paragraph_at(idx, "6.1.2. Hipótesis nula", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "La hipótesis nula establece que los algoritmos de Deep Learning no presentan un desempeño eficiente en la identificación de variedades de papas nativas cuando el promedio de Accuracy es menor al valor de referencia definido para la investigación.")
idx = insert_paragraph_at(idx, "6.1.3. Prueba estadística paramétrica utilizada", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "Para compatibilizar el diseño experimental computacional con la recomendación metodológica de la asesora, se utiliza la prueba t de Student para una muestra. Esta prueba compara el promedio de Accuracy obtenido en varias ejecuciones completas del entrenamiento contra el valor mínimo esperado μ₀ = 0.75.")
idx = add_caption(idx, 21, "Criterio de eficiencia para la interpretación del Accuracy")
idx = insert_table_at(idx, [
    ["Accuracy", "Interpretación"],
    ["< 75 %", "No es eficiente"],
    [">= 75 %", "Eficiente"],
])
idx = add_note(idx, "Nota. Elaboración propia a partir del umbral de eficiencia definido para la contrastación de hipótesis.")
idx = add_formula_block(idx, "t = (X̄ − μ₀) / (S / √N)", [
    "X̄ = promedio de Accuracy obtenido en las ejecuciones completas.",
    "μ₀ = valor esperado o umbral mínimo de eficiencia, igual a 0.75.",
    "S = desviación estándar de los valores de Accuracy.",
    "N = número de ejecuciones completas del entrenamiento.",
    "α = nivel de significancia, igual a 0.05.",
])
idx = insert_paragraph_at(idx, "Regla de decisión", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "Si p < 0.05, se rechaza H0. Si p ≥ 0.05, no se rechaza H0.", first_line=False)
idx = insert_paragraph_at(idx, "6.2. ANÁLISIS E INTERPRETACIÓN", heading=True, first_line=False)

# HG.
idx = insert_paragraph_at(idx, "6.2.1. Prueba de hipótesis general", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "HG: Los algoritmos de Deep Learning presentan un desempeño eficiente en la identificación de variedades de papas nativas en la comunidad de Santa Carmen.")
idx = insert_paragraph_at(idx, "H0G: Los algoritmos de Deep Learning no presentan un desempeño eficiente en la identificación de variedades de papas nativas en la comunidad de Santa Carmen.")
idx = insert_paragraph_at(idx, "H1G: Los algoritmos de Deep Learning presentan un desempeño eficiente en la identificación de variedades de papas nativas en la comunidad de Santa Carmen.")
idx = add_caption(idx, 22, "Resumen de ejecuciones para la hipótesis general")
idx = insert_table_at(idx, placeholder_execution_table("general", combined=True))
idx = add_note(idx, "Nota. Esta tabla se completa con los valores de Accuracy generados por los notebooks experimentales de CNN/EfficientNetB0 y Swin-Tiny Transformer.")
idx = add_formula_block(idx, "t = (X̄ − μ₀) / (S / √N)", ["X̄ = promedio general de Accuracy.", "μ₀ = 0.75.", "S = desviación estándar muestral.", "N = número de ejecuciones consideradas."])
idx = add_caption(idx, 23, "Prueba t de Student para una muestra en la hipótesis general")
idx = insert_table_at(idx, one_sample_t_table())
idx = add_note(idx)
idx = insert_paragraph_at(idx, "Conclusión de la hipótesis general", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "La decisión final se registrará después de completar las ejecuciones experimentales. Si p < 0.05 y el promedio de Accuracy es igual o superior a 0.75, se rechazará H0G y se aceptará la hipótesis general.")

# HE1.
idx = insert_paragraph_at(idx, "6.2.2. Prueba de hipótesis específica 1", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "HE1: El modelo CNN/EfficientNetB0 presenta un desempeño eficiente en la identificación de variedades de papas nativas.")
idx = insert_paragraph_at(idx, "H01: El modelo CNN/EfficientNetB0 presenta un desempeño ineficiente en la identificación de variedades de papas nativas.")
idx = insert_paragraph_at(idx, "H11: El modelo CNN/EfficientNetB0 presenta un desempeño eficiente en la identificación de variedades de papas nativas.")
idx = add_caption(idx, 24, "Ejecuciones del modelo CNN/EfficientNetB0 para la prueba de eficiencia")
idx = insert_table_at(idx, placeholder_execution_table("cnn_50_epocas"))
idx = add_note(idx, "Nota. Esta tabla se completa con `tabla_ejecuciones_accuracy_cnn.csv` generado por el notebook experimental de CNN/EfficientNetB0.")
idx = add_formula_block(idx, "t = (X̄ − μ₀) / (S / √N)", ["X̄ = promedio de Accuracy del modelo CNN/EfficientNetB0.", "μ₀ = 0.75.", "S = desviación estándar del Accuracy en las ejecuciones CNN.", "N = número de ejecuciones CNN."])
idx = add_caption(idx, 25, "Prueba t de Student para una muestra en el modelo CNN/EfficientNetB0")
idx = insert_table_at(idx, one_sample_t_table())
idx = add_note(idx)
idx = insert_paragraph_at(idx, "Conclusión de la hipótesis específica 1", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "Si p < 0.05 y el promedio de Accuracy del modelo CNN/EfficientNetB0 es igual o superior a 0.75, se rechazará H01 y se aceptará HE1, concluyendo que el modelo CNN presenta desempeño eficiente.")

# HE2.
idx = insert_paragraph_at(idx, "6.2.3. Prueba de hipótesis específica 2", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "HE2: El modelo Swin-Tiny Transformer presenta un desempeño eficiente en la identificación de variedades de papas nativas.")
idx = insert_paragraph_at(idx, "H02: El modelo Swin-Tiny Transformer presenta un desempeño ineficiente en la identificación de variedades de papas nativas.")
idx = insert_paragraph_at(idx, "H12: El modelo Swin-Tiny Transformer presenta un desempeño eficiente en la identificación de variedades de papas nativas.")
idx = add_caption(idx, 26, "Ejecuciones del modelo Swin-Tiny Transformer para la prueba de eficiencia")
idx = insert_table_at(idx, placeholder_execution_table("swin_50_epocas"))
idx = add_note(idx, "Nota. Esta tabla se completa con `tabla_ejecuciones_accuracy_swin.csv` generado por el notebook experimental de Swin-Tiny Transformer.")
idx = add_formula_block(idx, "t = (X̄ − μ₀) / (S / √N)", ["X̄ = promedio de Accuracy del modelo Swin-Tiny Transformer.", "μ₀ = 0.75.", "S = desviación estándar del Accuracy en las ejecuciones Swin.", "N = número de ejecuciones Swin."])
idx = add_caption(idx, 27, "Prueba t de Student para una muestra en el modelo Swin-Tiny Transformer")
idx = insert_table_at(idx, one_sample_t_table())
idx = add_note(idx)
idx = insert_paragraph_at(idx, "Conclusión de la hipótesis específica 2", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "Si p < 0.05 y el promedio de Accuracy del modelo Swin-Tiny Transformer es igual o superior a 0.75, se rechazará H02 y se aceptará HE2, concluyendo que el modelo Swin-Tiny Transformer presenta desempeño eficiente.")

# HE3.
idx = insert_paragraph_at(idx, "6.2.4. Prueba de hipótesis específica 3", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "HE3: Existen diferencias significativas entre el desempeño de CNN/EfficientNetB0 y Swin-Tiny Transformer en la identificación de variedades de papas nativas.")
idx = insert_paragraph_at(idx, "H03: No existen diferencias significativas entre el desempeño de CNN/EfficientNetB0 y Swin-Tiny Transformer en la identificación de variedades de papas nativas.")
idx = insert_paragraph_at(idx, "H13: Existen diferencias significativas entre el desempeño de CNN/EfficientNetB0 y Swin-Tiny Transformer en la identificación de variedades de papas nativas.")
idx = add_caption(idx, 28, "Comparación de Accuracy entre CNN/EfficientNetB0 y Swin-Tiny Transformer por ejecución")
idx = insert_table_at(idx, [
    ["Ejecución", "Seed", "Accuracy CNN", "Accuracy Swin", "Diferencia"],
    ["1", "Por completar", "Por completar", "Por completar", "Por completar"],
    ["2", "Por completar", "Por completar", "Por completar", "Por completar"],
    ["3", "Por completar", "Por completar", "Por completar", "Por completar"],
    ["4", "Por completar", "Por completar", "Por completar", "Por completar"],
    ["5", "Por completar", "Por completar", "Por completar", "Por completar"],
])
idx = add_note(idx, "Nota. Esta tabla se completa con ejecuciones equivalentes de ambos modelos. Cuando se usen las mismas semillas, la comparación puede tratarse como pareada.")
idx = add_formula_block(idx, "t = D̄ / (S_D / √N)", ["D̄ = promedio de las diferencias de Accuracy entre ambos modelos.", "S_D = desviación estándar de las diferencias.", "N = número de ejecuciones pareadas.", "α = 0.05."])
idx = add_caption(idx, 29, "Prueba t para la comparación entre CNN/EfficientNetB0 y Swin-Tiny Transformer")
idx = insert_table_at(idx, paired_t_table())
idx = add_note(idx)
idx = insert_paragraph_at(idx, "Conclusión de la hipótesis específica 3", heading=True, first_line=False)
idx = insert_paragraph_at(idx, "Si p < 0.05, se rechazará H03 y se aceptará HE3, concluyendo que existe diferencia significativa entre ambos modelos. Si p ≥ 0.05, no se rechazará H03 y la diferencia observada se interpretará como descriptiva, no estadísticamente significativa.")

# Ajustar captions de anexos porque las nuevas tablas ocupan 21-29.
for p in doc.paragraphs:
    text = norm(p.text)
    if text.startswith("Tabla 22") and "Cronograma" in text:
        p.text = "Tabla 30  | Cronograma de Actividades de investigación"
        format_paragraph(p, heading=True, first_line=False)
    elif text.startswith("Tabla 23") and "Presupuesto" in text:
        p.text = "Tabla 31  | Presupuesto del proyecto de investigación"
        format_paragraph(p, heading=True, first_line=False)

# Actualizar lista manual de tablas.
for i, p in enumerate(doc.paragraphs):
    text = norm(p.text)
    if text.startswith("Tabla 21") and "Matriz de contrastación" in text:
        p.text = "Tabla 21 Criterio de eficiencia para la interpretación del Accuracy 96"
    elif text.startswith("Tabla 22") and "Cronograma" in text:
        p.text = "Tabla 22 Resumen de ejecuciones para la hipótesis general 97"
    elif text.startswith("Tabla 23") and "Presupuesto" in text:
        p.text = "Tabla 23 Prueba t de Student para una muestra en la hipótesis general 98"
    if 0 <= i < 140 and p.text.startswith("Tabla "):
        format_paragraph(p, first_line=False)

# Insertar entradas faltantes después de Tabla 23 del índice, antes de RESUMEN.
body = doc._body._element
resume_idx = find_body_index_by_text("RESUMEN")
if resume_idx is not None:
    entries = [
        "Tabla 24 Ejecuciones del modelo CNN/EfficientNetB0 para la prueba de eficiencia 99",
        "Tabla 25 Prueba t de Student para una muestra en el modelo CNN/EfficientNetB0 100",
        "Tabla 26 Ejecuciones del modelo Swin-Tiny Transformer para la prueba de eficiencia 101",
        "Tabla 27 Prueba t de Student para una muestra en el modelo Swin-Tiny Transformer 102",
        "Tabla 28 Comparación de Accuracy entre CNN/EfficientNetB0 y Swin-Tiny Transformer por ejecución 103",
        "Tabla 29 Prueba t para la comparación entre CNN/EfficientNetB0 y Swin-Tiny Transformer 104",
        "Tabla 30 Cronograma de Actividades de investigación 109",
        "Tabla 31 Presupuesto del proyecto de investigación 110",
    ]
    for entry in entries:
        resume_idx = insert_paragraph_at(resume_idx, entry, first_line=False)

# Validación final.
doc.save(OUT)
check = Document(OUT)
all_text = "\n".join(p.text for p in check.paragraphs)
for table in check.tables:
    for row in table.rows:
        for cell in row.cells:
            all_text += "\n" + cell.text
forbidden = ["INKAPA", "aplicativo", "sintética", "sintético", "synthetic", "imágenes reales", "imágenes sintéticas", "etiqueta real", "Llunchuy Waqachiq"]
found = [term for term in forbidden if term.lower() in all_text.lower()]
print(f"Generado: {OUT}")
print(f"Párrafos: {len(check.paragraphs)} | Tablas: {len(check.tables)} | Imágenes internas: {len(check.inline_shapes)}")
print(f"Términos restringidos: {found}")
if found:
    raise SystemExit(2)
